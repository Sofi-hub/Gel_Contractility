"""
scripts/analyze_contractions.py
--------------------------------
Detección de contracciones sobre la salida de main.py.

Salidas -> data/processed_data/<nombre_video>/
    eventos.xlsx                  (hojas: eventos, resumen, por_segmento,
                                   perfil_frecuencia, estabilidad_umbral)
    02_eventos_detectados.png
    03_amplitudes.png
    04_perfil_frecuencia.png
    05_estabilidad_umbral.png
    06_comparacion_tramos.png     (única figura multipanel, a propósito)

Uso:
    python scripts/analyze_contractions.py --input data/processed_data/Video_prueba/serie_temporal.xlsx

VERIFICAR SIEMPRE en un video nuevo: mirar 05_estabilidad_umbral.png.
Si hay meseta, los eventos son reales y el conteo no depende del k
elegido. Si decae sin meseta, se está contando ruido.
"""

from __future__ import annotations
import argparse, sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from src import event_detection as ed
from src import plotting
from src.output_paths import video_output_dir
from src.qc_visualization import _write_xlsx


def _append_sheet(path, df, name):
    wb = load_workbook(path)
    ws = wb.create_sheet(name)
    HF, HFill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4472C4")
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=j, value=str(c)); cell.font = HF; cell.fill = HFill
        ws.column_dimensions[cell.column_letter].width = max(14, len(str(c)) + 2)
    for i, (_, row) in enumerate(df.iterrows(), 2):
        for j, c in enumerate(df.columns, 1):
            v = row[c]
            if isinstance(v, (np.floating, np.integer)): v = v.item()
            if isinstance(v, float) and np.isnan(v): v = None
            ws.cell(row=i, column=j, value=v).font = Font(name="Arial")
    wb.save(path)


def parse_args():
    p = argparse.ArgumentParser(description="Detección de contracciones (escala-invariante)")
    p.add_argument("--input", required=True, help="xlsx/csv de serie temporal (salida de main.py)")
    p.add_argument("--video-name", default=None,
                   help="Nombre para la carpeta de salida. Por defecto se deduce del --input.")
    p.add_argument("--output-dir", default=None, help="Sobrescribe la carpeta de salida")
    p.add_argument("--amp-k", type=float, default=6.0,
                   help="Umbral = k * ruido. Verificar con la curva de estabilidad.")
    p.add_argument("--sharpness", type=float, default=1.30,
                   help="Agudeza mínima (escala-invariante). 0 la desactiva.")
    p.add_argument("--min-amplitude-px", type=float, default=None,
                   help="Piso absoluto de amplitud, para descartar eventos no visibles a ojo.")
    p.add_argument("--noise-px", type=float, default=None, help="Ruido del video de control")
    p.add_argument("--freq-window-s", type=float, default=8.0,
                   help="Ventana del perfil de frecuencia. Subir para estimulación lenta.")
    p.add_argument("--zoom", nargs=2, type=float, action="append", metavar=("T0", "T1"),
                   help="Ventana de comparación, repetible. Si se omite, una por segmento.")
    return p.parse_args()


def main():
    a = parse_args()
    src = Path(a.input)
    df = pd.read_csv(src) if src.suffix == ".csv" else pd.read_excel(src, sheet_name="diagnostics")

    name = a.video_name or (src.parent.name if src.parent.name not in ("", ".") else src.stem)
    out = Path(a.output_dir) if a.output_dir else video_output_dir(name)
    print(f"Salida: {out}")

    kw = dict(sharpness_threshold=a.sharpness, min_amplitude_px=a.min_amplitude_px,
              noise_px=a.noise_px)
    r = ed.detect_contractions(df, amp_k=a.amp_k, **kw)

    print(f"Ancho de evento medido en los datos: {r.event_width_s:.3f} s")
    print(f"Ruido {'medido' if a.noise_px else 'estimado'}: {r.noise_px:.4f} px "
          f"| umbral: {r.amp_threshold_px:.4f} px")
    print(f"Candidatos: {r.diagnostics.get('n_candidatos',0)} | "
          f"descartados por agudeza: {r.diagnostics.get('descartados_por_agudeza',0)} | "
          f"EVENTOS: {len(r.events)}")

    scan = ed.threshold_stability_scan(df, **kw)
    fp = ed.symmetric_false_positive_check(df, amp_k=a.amp_k, **kw)
    print(f"Control de falsos positivos: {fp['n_abajo']} abajo / {fp['n_arriba']} arriba "
          f"-> {fp['veredicto']}")

    plotting.plot_threshold_stability(scan, out, amp_k_used=a.amp_k)

    if len(r.events) == 0:
        print("\n>>> NO se detectaron contracciones en este video.")
        print("    Revisá 05_estabilidad_umbral.png; si no hay meseta, no hay señal real.")
        return

    ev = ed.segment_by_rhythm(r.events)
    tt = df["time_s"].to_numpy()
    seg = ed.analyze_segments(ev, r.signal, tt)
    freq = ed.frequency_profile(r.signal, tt, window_s=a.freq_window_s)

    plotting.plot_events(df, r, ev, out)
    plotting.plot_amplitudes(ev, r, out)
    plotting.plot_frequency_profile(freq, out, window_s=a.freq_window_s)
    windows = [tuple(z) for z in a.zoom] if a.zoom else plotting.auto_windows_from_segments(seg)
    plotting.plot_segment_comparison(df, r, ev, out, windows)

    xlsx = out / "eventos.xlsx"
    _write_xlsx(ev, xlsx, summary={
        "archivo_fuente": str(src),
        "ancho_evento_medido_s": round(r.event_width_s, 4),
        "ruido_px": round(r.noise_px, 5),
        "ruido_es_provisorio": a.noise_px is None,
        "umbral_amplitud_px": round(r.amp_threshold_px, 5),
        "amp_k": a.amp_k,
        "umbral_agudeza": a.sharpness,
        "n_eventos": len(ev),
        "n_segmentos": int(ev.segmento.nunique()),
        "control_FP_abajo": fp["n_abajo"],
        "control_FP_arriba": fp["n_arriba"],
        "control_FP_veredicto": fp["veredicto"],
    })
    _append_sheet(xlsx, seg, "por_segmento")
    _append_sheet(xlsx, freq, "perfil_frecuencia")
    _append_sheet(xlsx, scan, "estabilidad_umbral")
    print(f"Tabla: {xlsx}")

    print("\n=== SEGMENTOS ===")
    for _, s in seg.iterrows():
        print(f"  Seg {int(s.segmento)}: {int(s.n_eventos)} eventos | "
              f"t={s.t_inicio_s:.1f}-{s.t_fin_s:.1f}s | {s.frecuencia_Hz:.3f} Hz "
              f"(T={s.periodo_s:.3f}s) | amp={s.amplitud_media_px:.3f} px | "
              f"CV={s.CV_intervalo_pct:.1f}% ({ed.interpret_regularity(s.CV_intervalo_pct)})")


if __name__ == "__main__":
    main()
