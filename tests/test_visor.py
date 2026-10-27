"""visor.py (pestaña Resultados): la parte de datos, sin ventana.

    python tests/test_visor.py

Comprueba que lo que muestra sale de los Excel vigentes (no recalcula):
mismos conteos, ficha de una contraccion con los valores de la hoja cinetica.
"""
from __future__ import annotations

import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
import visor  # noqa: E402
import numpy as np  # noqa: E402

ok = True


def check(nombre, cond):
    global ok
    print(f"  {nombre}: {'OK' if cond else 'FALLA'}")
    ok &= bool(cond)


P = RAIZ / "data" / "processed_data"
cs = visor.listar_carpetas(P)
check("lista de resultados (sin carpetas que empiezan con _)",
      len(cs) >= 1 and all(not c.name.startswith("_") for c in cs))
if (P / "Video_prueba").is_dir():
    r = visor.leer_resultado(P / "Video_prueba")
    check("Video_prueba: 29 contracciones en la hoja cinetica", len(r.cin) == 29)
    check("serie para dibujar del largo del video", len(r.senal) == len(r.t) > 1000)
    i = int((r.cin["grupo"] == "estimulados").to_numpy().nonzero()[0][0])
    f = dict(visor.ficha_evento(r, i))
    check("ficha de una estimulada: tipo, desvio y amplitud de la hoja",
          f["Tipo"] == "estimulada" and "Desvío del estimulador" in f
          and f"{r.cin['amplitud_px'].iloc[i]:.2f} px" in f["Amplitud"])
    check("cache: la segunda lectura devuelve el mismo objeto",
          visor.leer_resultado(P / "Video_prueba") is r)
if (P / "Video_613").is_dir():
    r = visor.leer_resultado(P / "Video_613")
    f = dict(visor.ficha_evento(r, 0))
    check("613: las fichas dicen que es candidata (NO reportable)", "candidata" in f["Contracción"])

# ---- Comparar, marcas y ficha (2026-10-11): solo leen y dibujan ----
import shutil, tempfile  # noqa: E402
nombres = [n for n in ("Video_prueba", "Video_063_CTRL1_5V", "Video_613") if (P / n).is_dir()]
if len(nombres) >= 2:
    rs = [visor.leer_resultado_sin_cache(P / n) for n in nombres]
    filas = dict(visor.tabla_comparar(rs))
    check("tabla comparar: una columna por resultado",
          all(len(v) == len(rs) for v in filas.values()))
    if "Video_prueba" in nombres:
        rp = rs[nombres.index("Video_prueba")]
        check("tabla comparar: Video_prueba 29 contracciones, amplitud del resumen",
              filas["Contracciones"][0] == "29"
              and filas["Amplitud, mediana (% del grosor)"][0].startswith(f"{rp.fila['amplitud_pct']:.2f} %"))
        lag, y, n = visor.promedio_para_dibujar(rp)
        check("promedio para dibujar: los 6 estimulados, pico cerca de la amplitud informada",
              n == 6 and abs(float(np.nanmax(y)) - rp.fila["amplitud_pct"]) < 0.1 * rp.fila["amplitud_pct"])
    if "Video_613" in nombres:
        check("613 en la tabla: NO reportable y sin amplitud",
              filas["Reportable"][nombres.index("Video_613")] == "NO"
              and filas["Amplitud, mediana (% del grosor)"][nombres.index("Video_613")] == "—")
    with tempfile.TemporaryDirectory() as tmp:
        x = visor.guardar_tabla_comparar(rs, Path(tmp) / "c.xlsx")
        import pandas as pd
        check("guardar tabla comparar (Excel)", pd.read_excel(x).shape == (len(filas), len(rs) + 1))

if (P / "Video_prueba").is_dir():
    with tempfile.TemporaryDirectory() as tmp:
        c = Path(tmp) / "Video_prueba"
        shutil.copytree(P / "Video_prueba", c)
        antes = {p.name: p.read_bytes() for p in c.glob("*.xlsx")}
        r = visor.leer_resultado_sin_cache(c)
        visor.marcar_dudosa(r, 3)
        check("marca dudosa: se guarda aparte", visor.leer_marcas(r) == {"3": "dudosa"})
        check("marca dudosa: los Excel no se tocan",
              {p.name: p.read_bytes() for p in c.glob("*.xlsx")} == antes)
        check("marca dudosa: la ficha la menciona", "Marca manual" in visor.ficha_texto(r, 2))
        visor.marcar_dudosa(r, 3, dudosa=False)
        check("quitar la última marca borra el archivo", not visor.archivo_marcas(c).exists())
        png = visor.guardar_ficha_png(r, 0, c / "f.png")
        check("exportar ficha como imagen", png.stat().st_size > 10000)

# ---- pulsos individuales y selección (2026-10-15): media de valores del Excel ----
import openpyxl  # noqa: E402


def _cinetica_cruda(carpeta: Path):
    """La hoja cinetica leída directo con openpyxl (sin visor ni pandas)."""
    x = sorted(carpeta.glob("contracciones_*.xlsx"))[0]
    wb = openpyxl.load_workbook(x, read_only=True, data_only=True)
    hoja = next(wb[n] for n in wb.sheetnames if n == "cinetica" or n.startswith("cinetica_"))
    filas = list(hoja.iter_rows(values_only=True))
    cab = list(filas[0])
    return [dict(zip(cab, f)) for f in filas[1:]]


for nombre, primeros in (("Video_063_CTRL1_5V", [2, 3, 4]), ("Video_466_EXP5_FAPS4_40V", [1, 2, 3])):
    if not (P / nombre).is_dir():
        continue
    r = visor.leer_resultado(P / nombre)
    cruda = _cinetica_cruda(P / nombre)
    check(f"{nombre}: tabla de pulsos, una fila por contracción", len(visor.tabla_pulsos(r)) == len(cruda))
    sel = visor.seleccion_rapida(r, "primeros", 3)
    check(f"{nombre}: 3 primeros estimulados = {primeros}", sel == primeros)
    d = visor.media_seleccion(r, sel)
    amp = [f["amplitud_relativa_pct"] for f in cruda if int(f["evento"]) in sel]
    px_ = [f["amplitud_px"] for f in cruda if int(f["evento"]) in sel]
    check(f"{nombre}: media de la selección = media de la hoja cinetica (% y px)",
          d["n"] == 3 and abs(d["amplitud_pct"] - float(np.mean(amp))) < 1e-9
          and abs(d["amplitud_px"] - float(np.mean(px_))) < 1e-9)
    ttp = [f["ttp_s"] for f in cruda if int(f["evento"]) in sel and f["ttp_medible"]]
    check(f"{nombre}: TTP de la selección = media de los medibles (o cota si ninguno)",
          (len(ttp) == 0 and d["ttp_n_medibles"] == 0 and d["ttp_cota_s"] > 0)
          or abs(d["ttp_s"] - float(np.mean(ttp))) < 1e-9)
    todos_est = visor.seleccion_rapida(r, "estimulados")
    dm = [f["amplitud_relativa_pct"] for f in cruda if int(f["evento"]) in todos_est]
    check(f"{nombre}: la MEDIANA de todos los estimulados es la amplitud informada",
          abs(float(np.median(dm)) - r.fila["amplitud_pct"]) < 1e-9)
    check(f"{nombre}: rótulo 'media de la selección: 3 de N pulsos'",
          visor.texto_media(d)[0][1].startswith(f"3 de {len(cruda)} pulsos"))
    check(f"{nombre}: ninguno y todos", visor.seleccion_rapida(r, "ninguno") == []
          and len(visor.seleccion_rapida(r, "todos")) == len(cruda))

if (P / "Video_613").is_dir():
    r = visor.leer_resultado(P / "Video_613")
    check("613 (NO reportable): sin estimulados, 'primeros' toma los primeros candidatos",
          visor.seleccion_rapida(r, "estimulados") == [] and len(visor.seleccion_rapida(r, "primeros", 3)) == 3
          and all(f["tipo"] == "candidata" for f in visor.tabla_pulsos(r)))

if (P / "Video_063_CTRL1_5V").is_dir():
    with tempfile.TemporaryDirectory() as tmp:
        c = Path(tmp) / "Video_063_CTRL1_5V"
        shutil.copytree(P / "Video_063_CTRL1_5V", c)
        antes = {p.name: p.read_bytes() for p in c.glob("*.xlsx")}
        r = visor.leer_resultado_sin_cache(c)
        check("sin selección guardada: None", visor.leer_seleccion(r) is None)
        visor.guardar_seleccion(r, [4, 2], "a mano")
        check("selección guardada aparte y releída", visor.leer_seleccion(r) == [2, 4])
        ev, txt = visor.eventos_para_comparar(r, "guardada")
        check("Comparar usa la selección guardada", ev == [2, 4] and "guardada" in txt)
        x = visor.exportar_seleccion(r, [2, 3], c / "sel.xlsx")
        import pandas as pd
        check("exportar selección: 2 pulsos y la media",
              len(pd.read_excel(x, sheet_name="pulsos_elegidos")) == 2
              and "Media de la selección" in pd.read_excel(x, sheet_name="media")["qué"].tolist())
        x2 = visor.exportar_seleccion(r, [2, 3], c / "sel.csv")
        check("exportar selección como CSV", x2.stat().st_size > 50)
        check("selección: los Excel no se tocan",
              {p.name: p.read_bytes() for p in c.glob("*.xlsx") if p.name != "sel.xlsx"} == antes)

nombres = [n for n in ("Video_063_CTRL1_5V", "Video_466_EXP5_FAPS4_40V") if (P / n).is_dir()]
if len(nombres) == 2:
    rs = [visor.leer_resultado(P / n) for n in nombres]
    filas = dict(visor.tabla_seleccion(rs, "primeros", 3))
    check("Comparar: media de los 3 primeros de cada uno",
          filas["Media: amplitud (% del grosor)"] ==
          [f"{visor.media_seleccion(r, visor.seleccion_rapida(r, 'primeros', 3))['amplitud_pct']:.2f} %"
           for r in rs])
    lag, ys, ym = visor.pulsos_para_dibujar(rs[1], [1, 2, 3])
    check("dibujo de la selección: 3 pulsos y su media", len(ys) == 3 and ym.shape == lag.shape)

# ---- Comparar: pulsos elegidos y vistas (2026-10-15) ----
if len(nombres) == 2:
    df = visor.pulsos_comparar(rs, "primeros", 3)
    check("Comparar, pulsos elegidos: 3 por video, con su resultado",
          len(df) == 6 and df["resultado"].str.startswith("1 · ").sum() == 3)
    v = visor.valores_pulsos(rs[1], [3, 1, 2])
    check("valores de los pulsos en orden de tiempo", v["evento"].tolist() == [1, 2, 3])
check("compactar números de pulso", visor.compactar([1, 2, 3, 5, 7, 8]) == "1–3, 5, 7, 8")

# ---- marcas y selección por INSTANTE (2026-10-26) ----
# Si se rehace el paso 2 y cambia la numeración, la marca tiene que seguir en la
# misma contracción; y si esa contracción ya no está, no se marca otra.
import json  # noqa: E402
import pandas as pd  # noqa: E402


def _falso(carpeta, eventos):
    """Un Resultado mínimo: solo lo que miran las marcas y la selección (la
    carpeta y, de la hoja cinetica, el número y el instante de cada una)."""
    return visor.Resultado(carpeta=Path(carpeta),
                           cin=pd.DataFrame({"evento": [n for n, _ in eventos],
                                             "tiempo_s": [t for _, t in eventos]}))


with tempfile.TemporaryDirectory() as tmp:
    c = Path(tmp)
    r_antes = _falso(c, [(1, 5.00), (2, 12.30), (3, 20.10)])
    visor.marcar_dudosa(r_antes, 2)
    check("marca: el archivo guarda el instante (y el número al lado)",
          json.loads(visor.archivo_marcas(c).read_text(encoding="utf-8"))["dudosas"]
          == [{"instante_s": 12.3, "n": 2, "nota": "dudosa"}])
    # paso 2 rehecho: aparece una contracción nueva al principio y todo se renumera
    r_despues = _falso(c, [(1, 2.00), (2, 5.00), (3, 12.30), (4, 20.10)])
    m, av = visor.marcas_y_avisos(r_despues)
    check("marca: tras renumerar sigue en la misma contracción (2 -> 3)",
          m == {"3": "dudosa"} and av == [])
    # y si la contracción marcada ya no está, no se marca otra: se avisa
    m2, av2 = visor.marcas_y_avisos(_falso(c, [(1, 5.00), (2, 20.10)]))
    check("marca sin contracción: aviso y ninguna otra marcada",
          m2 == {} and len(av2) == 1 and av2[0].startswith("marca de 12.30 s sin contracción"))

    visor.guardar_seleccion(r_antes, [1, 3], "a mano")
    check("selección: tras renumerar son las mismas contracciones (1, 3 -> 2, 4)",
          visor.leer_seleccion(r_despues) == [2, 4])
    ev, av3 = visor.seleccion_y_avisos(_falso(c, [(1, 5.00)]))
    check("pulso elegido sin contracción: aviso y no se elige otro",
          ev == [1] and len(av3) == 1 and av3[0].startswith("pulso elegido de 20.10 s sin contracción"))

    # archivos del formato viejo (solo el número): se siguen leyendo
    visor.archivo_marcas(c).write_text(json.dumps({"dudosas": {"2": "dudosa"}}), encoding="utf-8")
    visor.archivo_seleccion(c).write_text(json.dumps({"eventos": [1, 3]}), encoding="utf-8")
    check("archivo viejo (solo el número): marcas y selección se leen igual",
          visor.leer_marcas(r_antes) == {"2": "dudosa"} and visor.leer_seleccion(r_antes) == [1, 3])
    visor.marcar_dudosa(r_antes, 3)
    check("al guardar, un archivo viejo queda en el formato nuevo",
          json.loads(visor.archivo_marcas(c).read_text(encoding="utf-8"))["dudosas"]
          == [{"instante_s": 12.3, "n": 2, "nota": "dudosa"},
              {"instante_s": 20.1, "n": 3, "nota": "dudosa"}])
    visor.marcar_dudosa(r_antes, 2, dudosa=False)
    visor.marcar_dudosa(r_antes, 3, dudosa=False)
    check("quitar la última marca borra el archivo", not visor.archivo_marcas(c).exists())


# ---- basal (C8): el texto sale del resumen; resultados viejos lo avisan ----
import informe  # noqa: E402
check("basal: resultado viejo -> avisa que falta el paso 2",
      "rehacer el paso 2" in informe.texto_basal({"amplitud_relativa_pct": 1.0}))
check("basal: texto guardado", informe.texto_basal({"basal_resumen": "0 contraccion(es)", "basal_fin_s": 10.0})
      == "0–10.0 s: 0 contraccion(es)")

# ---- respuesta forzada (2026-10-19): texto del resumen; sin calcular nada ----
check("forzada: sin columnas -> no se muestra", informe.texto_forzada({"amplitud_relativa_pct": 1.0}) is None)
_rf = {"forzada_p": 0.024, "forzada_periodo_s": 10.1, "forzada_hay_respuesta": True,
       "forzada_n_responden": 6, "forzada_n_pulsos": 6, "forzada_amplitud_pct": 0.262,
       "forzada_amplitud_px": 6.02}
check("forzada: hay respuesta, con amplitud y pulsos",
      informe.texto_forzada(_rf).startswith("HAY") and "6.02 px" in informe.texto_forzada(_rf)
      and "6 de 6" in informe.texto_forzada(_rf))
check("forzada: sin respuesta",
      informe.texto_forzada({**_rf, "forzada_hay_respuesta": False, "forzada_p": 0.78}).startswith("sin respuesta"))
_cl, _fr = informe.frase_resultado({"fila": {"eventos": 8, "reportable": False}, "res": _rf})
check("forzada: la frase del informe la menciona si el conteo no es reportable",
      _cl == "mal" and "respuesta forzada" in _fr and "0.262" in _fr)
sys.path.insert(0, str(RAIZ / "scripts"))
import procesar_carpeta as _pc  # noqa: E402
check("forzada: tabla de carpeta",
      _pc._txt_forzada({"forzada_p": 0.024, "forzada_amplitud_px": 6.02, "forzada_pulsos": "6/6"}) == "6.02 px 6/6"
      and _pc._txt_forzada({"forzada_p": 0.78, "forzada_amplitud_px": None}) == "no"
      and _pc._txt_forzada({}) == "")

print("\nTODO OK" if ok else "\nHAY FALLAS")
sys.exit(0 if ok else 1)
