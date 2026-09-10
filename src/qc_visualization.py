"""
qc_visualization.py
--------------------
Control de calidad visual: dibuja sobre un frame los bordes crudos
detectados por columna, marca cuáles fueron aceptados como "inlier" y
cuáles descartados, y superpone el modelo ajustado final. También
incluye una vista de "microscopio" de UNA columna (perfil de intensidad
+ gradiente).

NUEVO: `plot_roi_profile()` grafica el perfil de grosor y de nitidez de
TODA la imagen con la ROI elegida marcada encima. Ése es el gráfico que
responde "¿por qué la ROI se cortó en x=577 y no siguió?" — sin él, la
selección de gauge region es una caja negra.
"""

from __future__ import annotations
from dataclasses import dataclass
import numpy as np
import pandas as pd
import cv2
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from . import edge_detection
from . import robust_fitting
from .pipeline import PipelineConfig


@dataclass
class FrameDiagnostics:
    """Todo lo que produce el pipeline en un frame, sin colapsar a un
    único número de grosor."""
    x: np.ndarray
    y_top: np.ndarray
    y_bottom: np.ndarray
    quality: np.ndarray
    top_fit: robust_fitting.RobustEdgeFit | None
    bottom_fit: robust_fitting.RobustEdgeFit | None
    thickness_px: float


def run_frame_diagnostics(
    frame_gray: np.ndarray,
    x_positions: np.ndarray,
    top_guess: np.ndarray,
    bottom_guess: np.ndarray,
    config: PipelineConfig,
) -> FrameDiagnostics:
    """
    Corre el mismo camino que pipeline.process_frame, pero sin colapsar
    el resultado: devuelve los puntos crudos por columna y los objetos de
    ajuste completos (con la máscara de inliers).
    """
    x, y_top, y_bot, quality = edge_detection.extract_edges_for_frame(
        frame_gray, x_positions, top_guess, bottom_guess,
        half_window=config.half_window, method=config.edge_method,
        min_gradient=config.min_gradient,
    )

    if config.fit_method == "ransac":
        kw = dict(degree=config.ransac_degree,
                  residual_threshold=config.ransac_residual_threshold,
                  residual_k=config.ransac_residual_k,
                  residual_floor=config.ransac_residual_floor)
        top_fit = robust_fitting.fit_edge_ransac(x, y_top, **kw)
        bot_fit = robust_fitting.fit_edge_ransac(x, y_bot, **kw)
    else:
        top_fit = robust_fitting.fit_edge_median(x, y_top)
        bot_fit = robust_fitting.fit_edge_median(x, y_bot)

    if top_fit is not None and bot_fit is not None:
        thickness_px = float(np.median(bot_fit.y_fitted - top_fit.y_fitted))
    else:
        thickness_px = float("nan")

    return FrameDiagnostics(
        x=x, y_top=y_top, y_bottom=y_bot, quality=quality,
        top_fit=top_fit, bottom_fit=bot_fit, thickness_px=thickness_px,
    )


def diagnostics_to_dataframe(diag: FrameDiagnostics) -> pd.DataFrame:
    """Tabla columna-por-columna, con el residuo de cada punto contra el
    modelo ajustado — que es lo que RANSAC realmente compara contra el
    umbral. Si ves residuos grandes y CONTIGUOS, el problema es el modelo
    (curvatura), no burbujas: las burbujas dan residuos aislados."""
    n = len(diag.x)
    top_inlier = diag.top_fit.inlier_mask if diag.top_fit is not None else np.zeros(n, dtype=bool)
    bot_inlier = diag.bottom_fit.inlier_mask if diag.bottom_fit is not None else np.zeros(n, dtype=bool)
    top_res = (diag.y_top - diag.top_fit.y_fitted) if diag.top_fit is not None else np.full(n, np.nan)
    bot_res = (diag.y_bottom - diag.bottom_fit.y_fitted) if diag.bottom_fit is not None else np.full(n, np.nan)
    return pd.DataFrame({
        "x": diag.x,
        "y_top_raw": diag.y_top,
        "y_bottom_raw": diag.y_bottom,
        "gradient_quality": diag.quality,
        "top_is_inlier": top_inlier,
        "bottom_is_inlier": bot_inlier,
        "top_residual_px": np.round(top_res, 4),
        "bottom_residual_px": np.round(bot_res, 4),
        "thickness_raw": diag.y_bottom - diag.y_top,
    })


def draw_diagnostics_overlay(
    frame_gray: np.ndarray,
    diag: FrameDiagnostics,
    px_to_mm: float = 1.0,
) -> np.ndarray:
    """
    Imagen BGR con:
        - VERDE: borde crudo aceptado como inlier.
        - ROJO: borde crudo descartado como outlier (posible burbuja).
        - CRUZ AMARILLA: columna sin borde confiable.
        - CIAN: modelo robusto final (lo que define el grosor).
    """
    img = cv2.cvtColor(frame_gray, cv2.COLOR_GRAY2BGR)

    GREEN, RED, YELLOW, CYAN, WHITE = (0, 220, 0), (0, 0, 255), (0, 230, 230), (255, 255, 0), (255, 255, 255)

    top_guess_fallback = int(np.nanmedian(diag.y_top)) if not np.all(np.isnan(diag.y_top)) else 0

    def draw_points(y_arr, inlier_mask):
        for i, (xi, yi) in enumerate(zip(diag.x, y_arr)):
            if np.isnan(yi):
                cv2.drawMarker(img, (int(xi), top_guess_fallback), YELLOW,
                               markerType=cv2.MARKER_TILTED_CROSS, markerSize=8, thickness=1)
                continue
            color = GREEN if (inlier_mask is not None and inlier_mask[i]) else RED
            cv2.circle(img, (int(xi), int(round(yi))), 2, color, thickness=-1)

    top_inlier = diag.top_fit.inlier_mask if diag.top_fit is not None else None
    bot_inlier = diag.bottom_fit.inlier_mask if diag.bottom_fit is not None else None

    draw_points(diag.y_top, top_inlier)
    draw_points(diag.y_bottom, bot_inlier)

    for fit in (diag.top_fit, diag.bottom_fit):
        if fit is not None:
            pts = np.stack([diag.x, fit.y_fitted], axis=1).astype(np.int32)
            cv2.polylines(img, [pts], isClosed=False, color=CYAN, thickness=1, lineType=cv2.LINE_AA)

    thickness_mm = diag.thickness_px * px_to_mm
    n_out = sum(f.n_outliers for f in (diag.top_fit, diag.bottom_fit) if f is not None)
    unit = "mm" if abs(px_to_mm - 1.0) > 1e-9 else "px SIN CALIBRAR"
    label = f"grosor={diag.thickness_px:.2f}px ({thickness_mm:.4f} {unit})  outliers={n_out}/{2*len(diag.x)}"
    cv2.putText(img, label, (10, 25), cv2.FONT_HERSHEY_SIMPLEX, 0.6, WHITE, 1, cv2.LINE_AA)

    # Segunda línea: umbral y residuo efectivos. Si el residuo es del
    # mismo orden que el umbral, RANSAC está separando ruido de ruido.
    if diag.top_fit is not None and diag.bottom_fit is not None:
        info = (f"umbral RANSAC: sup={diag.top_fit.threshold_px:.2f}px "
                f"inf={diag.bottom_fit.threshold_px:.2f}px  |  "
                f"residuo (MAD): sup={diag.top_fit.residual_px:.2f}px "
                f"inf={diag.bottom_fit.residual_px:.2f}px")
        cv2.putText(img, info, (10, 48), cv2.FONT_HERSHEY_SIMPLEX, 0.5, WHITE, 1, cv2.LINE_AA)

    cv2.putText(img, "verde=inlier  rojo=outlier(burbuja?)  amarillo=sin borde  cian=ajuste final",
                (10, img.shape[0] - 12), cv2.FONT_HERSHEY_SIMPLEX, 0.45, WHITE, 1, cv2.LINE_AA)

    return img


def plot_roi_profile(roi: dict, output_path=None):
    """
    Perfil de grosor y de nitidez columna por columna, con la ROI elegida
    sombreada. Responde de un vistazo:

      - ¿la ROI cayó sobre una zona PLANA (gauge region) o sobre el
        hombro del anclaje?
      - ¿por qué se cortó donde se cortó: cayó la nitidez, se disparó el
        grosor, o se perdió el seguimiento de la franja?

    Devuelve la figura de matplotlib.
    """
    T = roi.get("thickness_profile")
    S = roi.get("sharpness_profile")
    valid = roi.get("valid_columns")
    if T is None:
        raise ValueError("El dict de ROI no trae thickness_profile (¿versión vieja de auto_detect_roi?)")

    x = np.arange(len(T))
    q = roi.get("roi_quality", {})

    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(12, 6), sharex=True)

    ax1.plot(x, T, color="steelblue", lw=1.2, label="grosor (suavizado)")
    if valid is not None:
        ax1.plot(x[~valid], np.full((~valid).sum(), np.nanmin(T)), "|", color="darkred",
                 ms=6, label="franja no seguida")
    if q.get("cintura_px"):
        ax1.axhline(q["cintura_px"], color="gray", ls=":", label=f"cintura = {q['cintura_px']} px")
    ax1.axvspan(roi["x_start"], roi["x_end"], color="mediumseagreen", alpha=0.22,
                label=f"ROI usada ({roi['x_start']}-{roi['x_end']})")
    ax1.set_ylabel("Grosor (px)")
    ax1.set_title(f"Perfil del gel y ROI elegida — metodo: {q.get('method')} "
                  f"| variacion dentro de la ROI: {q.get('variacion_en_roi_pct')}%")
    ax1.legend(fontsize=8, ncol=2)
    ax1.grid(alpha=0.3)

    ax2.plot(x, S, color="darkorange", lw=1.0, label="nitidez del borde (min de los dos)")
    ax2.axvspan(roi["x_start"], roi["x_end"], color="mediumseagreen", alpha=0.22)
    ax2.set_ylabel("Gradiente vertical")
    ax2.set_xlabel("Columna x (px)")
    ax2.legend(fontsize=8)
    ax2.grid(alpha=0.3)

    fig.tight_layout()
    if output_path is not None:
        fig.savefig(output_path, dpi=140)
    return fig


def save_diagnostics(df: pd.DataFrame, output_path, fmt: str = "xlsx", summary: dict | None = None):
    """
    Guarda la tabla de diagnóstico. Por defecto .xlsx nativo (Excel no
    muestra el aviso de "se pueden perder datos" de los .csv), con las
    filas outlier resaltadas y una hoja de resumen.
    """
    from pathlib import Path
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if fmt in ("csv", "both"):
        df.to_csv(output_path.with_suffix(".csv"), index=False)

    if fmt in ("xlsx", "both"):
        _write_xlsx(df, output_path.with_suffix(".xlsx"), summary)

    return output_path.with_suffix(".xlsx" if fmt != "csv" else ".csv")


def _write_xlsx(df: pd.DataFrame, path, summary: dict | None = None):
    """Escribe el xlsx con formato: encabezado fijo, autofiltro, ancho de
    columnas, y filas outlier resaltadas."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "diagnostics"

    HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill("solid", fgColor="4472C4")
    BODY_FONT = Font(name="Arial")
    OUTLIER_FILL = PatternFill("solid", fgColor="FCE4E4")
    MISSING_FILL = PatternFill("solid", fgColor="FFF2CC")

    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=j, value=str(col))
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")

    for i, (_, row) in enumerate(df.iterrows(), start=2):
        is_outlier = (not bool(row.get("top_is_inlier", True))) or (not bool(row.get("bottom_is_inlier", True)))
        is_missing = pd.isna(row.get("y_top_raw")) or pd.isna(row.get("y_bottom_raw"))

        for j, col in enumerate(df.columns, start=1):
            val = row[col]
            if isinstance(val, (np.floating, np.integer)):
                val = val.item()
            if isinstance(val, float) and np.isnan(val):
                val = None
            if isinstance(val, (np.bool_, bool)):
                val = bool(val)
            c = ws.cell(row=i, column=j, value=val)
            c.font = BODY_FONT
            if isinstance(val, float):
                c.number_format = "0.000"
            if is_missing:
                c.fill = MISSING_FILL
            elif is_outlier:
                c.fill = OUTLIER_FILL

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df)+1}"
    for j, col in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(14, len(str(col)) + 2)

    if summary:
        ws2 = wb.create_sheet("resumen")
        ws2.cell(row=1, column=1, value="Métrica").font = HEADER_FONT
        ws2.cell(row=1, column=1).fill = HEADER_FILL
        ws2.cell(row=1, column=2, value="Valor").font = HEADER_FONT
        ws2.cell(row=1, column=2).fill = HEADER_FILL
        for i, (k, v) in enumerate(summary.items(), start=2):
            ws2.cell(row=i, column=1, value=str(k)).font = BODY_FONT
            ws2.cell(row=i, column=2, value=v if isinstance(v, (int, float, str)) else str(v)).font = BODY_FONT
        ws2.column_dimensions["A"].width = 34
        ws2.column_dimensions["B"].width = 34

    wb.save(path)


def plot_column_profile(
    frame_gray: np.ndarray,
    x: int,
    guess_row: float,
    half_window: int,
    polarity: int,
    min_gradient: float = 5.0,
):
    """
    Vista "de microscopio" de UNA columna: perfil de intensidad, su
    gradiente, y el punto subpíxel detectado, con la línea de
    min_gradient superpuesta. Responde "¿por qué esta columna se
    descartó?".
    """
    h = frame_gray.shape[0]
    profile = frame_gray[:, x].astype(np.float64)

    start = max(0, int(guess_row - half_window))
    end = min(h, int(guess_row + half_window))

    window = profile[start:end]
    grad = np.gradient(window) * polarity
    y_axis_full = np.arange(start, end)

    pt = edge_detection.subpixel_edge_parabolic(profile, start, end, polarity=polarity,
                                                min_gradient=min_gradient)

    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(7, 6), sharex=True)

    ax1.plot(y_axis_full, window, marker="o", ms=3, color="steelblue")
    ax1.set_ylabel("Intensidad")
    ax1.set_title(f"Columna x={x} | polarity={polarity} | válido={pt.valid}")
    if pt.valid:
        ax1.axvline(pt.y, color="crimson", linestyle="--", label=f"borde subpíxel y={pt.y:.2f}")
        ax1.legend()

    ax2.plot(y_axis_full, grad, marker="o", ms=3, color="darkorange")
    ax2.axhline(min_gradient, color="gray", linestyle=":", label=f"min_gradient={min_gradient}")
    ax2.set_ylabel("Gradiente (dI/dy · polarity)")
    ax2.set_xlabel("Fila (y, píxeles)")
    if pt.valid:
        ax2.axvline(pt.y, color="crimson", linestyle="--")
    ax2.legend()

    fig.tight_layout()
    return fig
